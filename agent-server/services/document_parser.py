"""本地文档解析器：将常见格式转换为可检索的纯文本。"""

from __future__ import annotations

import io
import csv
import json
import re
from html.parser import HTMLParser
from pathlib import Path


class DocumentParseError(ValueError):
    """文档无法解析或没有可用文本。"""


def parse_document(filename: str, data: bytes) -> tuple[str, dict]:
    """按扩展名解析文档，返回正文和解析元数据。"""
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        text = _decode_text(data)
        metadata = {"format": "txt"}
    elif suffix in {".md", ".markdown"}:
        text, metadata = _parse_text_format(data, suffix[1:])
    elif suffix in {".csv", ".tsv"}:
        text, metadata = _parse_csv(data, suffix)
    elif suffix == ".json":
        text, metadata = _parse_json(data)
    elif suffix in {".html", ".htm", ".xhtml"}:
        text, metadata = _parse_html(data)
    elif suffix == ".pdf":
        text, metadata = _parse_pdf(data)
    elif suffix == ".docx":
        text, metadata = _parse_docx(data)
    elif suffix == ".doc":
        raise DocumentParseError("暂不支持旧版 .doc，请另存为 .docx 或 PDF 后上传")
    elif suffix == ".xls":
        raise DocumentParseError("暂不支持旧版 .xls，请另存为 .xlsx 后上传")
    elif suffix in {".xlsx", ".pptx"}:
        text, metadata = _parse_office_xml(data, suffix)
    elif suffix == ".ppt":
        raise DocumentParseError("暂不支持旧版 .ppt，请另存为 .pptx 后上传")
    else:
        raise DocumentParseError(f"不支持的文件格式: {suffix}")
    text = text.replace("\x00", "").strip()
    if not text:
        raise DocumentParseError("文档中没有可提取的文本；扫描 PDF 需要先进行 OCR")
    return text, metadata


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "big5"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("文本文件编码无法识别")


def _parse_text_format(data: bytes, format_name: str) -> tuple[str, dict]:
    return _decode_text(data), {"format": format_name}


def _parse_csv(data: bytes, suffix: str) -> tuple[str, dict]:
    text = _decode_text(data)
    delimiter = "\t" if suffix == ".tsv" else csv.Sniffer().sniff(text[:4096]).delimiter if text.strip() else ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    normalized = "\n".join(" | ".join(cell.strip() for cell in row) for row in rows)
    return normalized, {"format": suffix[1:], "rows": len(rows), "delimiter": delimiter}


def _parse_json(data: bytes) -> tuple[str, dict]:
    try:
        value = json.loads(_decode_text(data))
        return json.dumps(value, ensure_ascii=False, indent=2), {"format": "json"}
    except json.JSONDecodeError as exc:
        raise DocumentParseError(f"JSON 解析失败: {exc}") from exc


class _HTMLTextParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        value = data.strip()
        if value:
            self.parts.append(value)


def _parse_html(data: bytes) -> tuple[str, dict]:
    parser = _HTMLTextParser()
    parser.feed(_decode_text(data))
    return "\n".join(parser.parts), {"format": "html"}


def _parse_office_xml(data: bytes, suffix: str) -> tuple[str, dict]:
    try:
        if suffix == ".xlsx":
            from openpyxl import load_workbook
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            blocks = []
            for sheet in workbook.worksheets:
                blocks.append(f"[工作表: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                    if values:
                        blocks.append(" | ".join(values))
            return "\n".join(blocks), {"format": "xlsx", "sheets": len(workbook.worksheets)}
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(data))
        blocks = []
        for index, slide in enumerate(presentation.slides, 1):
            blocks.append(f"[第 {index} 页]")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    blocks.append(shape.text.strip())
        return "\n".join(blocks), {"format": "pptx", "slides": len(presentation.slides)}
    except ImportError as exc:
        dependency = "openpyxl" if suffix == ".xlsx" else "python-pptx"
        raise DocumentParseError(f"缺少 {suffix} 解析依赖，请安装 {dependency}") from exc
    except Exception as exc:
        raise DocumentParseError(f"{suffix} 解析失败: {exc}") from exc


def _parse_pdf(data: bytes) -> tuple[str, dict]:
    try:
        import fitz
        document = fitz.open(stream=data, filetype="pdf")
        pages = [page.get_text("text").strip() for page in document]
        text = "\n\n".join(f"[第 {i + 1} 页]\n{page}" for i, page in enumerate(pages) if page)
        return text, {"format": "pdf", "pages": len(pages), "extracted_pages": sum(bool(p) for p in pages)}
    except ImportError as exc:
        raise DocumentParseError("缺少 PDF 解析依赖 PyMuPDF，请安装 pymupdf") from exc
    except Exception as exc:
        raise DocumentParseError(f"PDF 解析失败: {exc}") from exc


def _parse_docx(data: bytes) -> tuple[str, dict]:
    try:
        from docx import Document
        document = Document(io.BytesIO(data))
        blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(" | ".join(cells))
        return "\n".join(blocks), {"format": "docx", "paragraphs": len(document.paragraphs), "tables": len(document.tables)}
    except ImportError as exc:
        raise DocumentParseError("缺少 DOCX 解析依赖 python-docx，请安装 python-docx") from exc
    except Exception as exc:
        raise DocumentParseError(f"DOCX 解析失败: {exc}") from exc
